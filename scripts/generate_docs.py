"""Generate 技术方案.docx and 使用说明书.docx from scratch."""

import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def _add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x24, 0x4A, 0x68)
    return h


def _add_body(doc, text):
    p = doc.add_paragraph(text)
    style = doc.styles["Normal"]
    if style.font.size is None:
        style.font.size = Pt(11)
    p.style = style
    return p


def _add_table(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            table.rows[ri + 1].cells[ci].text = str(cell)
    return table


# ============================================================
# 技术方案
# ============================================================


def generate_tech_proposal():
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    # Title
    title = doc.add_heading("He-3 核反应堆探测器脉冲信号模拟系统", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("技术方案").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"版本 0.1.0 | {datetime.date.today()}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    # 1. 项目概述
    _add_heading(doc, "一、项目概述")
    _add_body(
        doc,
        (
            "本项目旨在建立一个可运行、可测试、可扩展的科研软件工程，用于模拟零功率反应堆场景下 "
            "He-3 正比计数管及其前置放大器输出的脉冲信号。系统以死时间之前的真实中子到达率为主要"
            "工况输入，从统计正确的物理过程生成可复现的连续脉冲波形，并通过相关中子噪声分析方法"
            "（Rossi-α、Feynman-α、PSD、连续 ACF/VTM、Wiener 去卷积）实现 α 复原的闭环验证。"
        ),
    )
    _add_body(
        doc,
        (
            "项目采用数字孪生三级验证阶梯：Tier 1 软件孪生（in-silico，用已知真值自证正确）、"
            "Tier 2 硬件在环（DT5800 验证电子学链）、Tier 3 零功率反应堆（真实裂变链验证）。"
            "当前实现聚焦于 Tier 1，完成纯瞬发相关中子事件生成（Phase A）、脉冲模式噪声分析"
            "（Phase B）、连续信号噪声分析（Phase C），以及可交互 Streamlit 界面。"
        ),
    )

    # 2. 物理模型
    _add_heading(doc, "二、物理模型")
    _add_heading(doc, "2.1 中子到达过程", level=2)
    _add_body(
        doc,
        (
            "非相关中子采用恒定计数率齐次泊松过程。相关中子采用一速点模型纯瞬发分支过程："
            "外源泊松注入 → 每个中子经历指数寿命后发生裂变/俘获/探测反应。参数映射集中校验 "
            "k_eff、α、探测效率 ε、ν̄ 与源强 S 的自洽性。"
        ),
    )
    _add_body(doc, "核心参数关系：")
    _add_body(doc, "  λ_t = α / (1 − k_eff)")
    _add_body(doc, "  λ_f = k_eff · λ_t / ν̄")
    _add_body(doc, "  λ_d = ε · λ_t")
    _add_body(doc, "  R_d = S · ε / (1 − k_eff)    (一阶探测率)")

    _add_heading(doc, "2.2 He-3 沉积能谱", level=2)
    _add_body(
        doc,
        "参数化混合模型：764 keV 全能峰（截断高斯）+ "
        "质子壁效应连续区 + 氚壁效应连续区 + 可选双壁效应。"
        "混合权重非负且和为 1，所有形状参数为 synthetic_demo，预留 MCNP 导入接口。",
    )

    _add_heading(doc, "2.3 双指数脉冲模型", level=2)
    _add_body(
        doc,
        (
            "单脉冲采用峰值归一化双指数：v(t) = polarity · A_peak · h(t − t_i)，"
            "其中 h(s) = g(s)/g(t_peak)，g(s) = exp(−s/τ_d) − exp(−s/τ_r)，τ_d > τ_r > 0。"
            "离散实现按事件分数相位校正峰值，保证 A_peak 与目标值一致。"
        ),
    )

    _add_heading(doc, "2.4 连续波形与信号处理", level=2)
    _add_body(
        doc,
        (
            "连续模拟电压 v(t) = b(t) + n(t) + Σ_i v_i(t)，其中 b(t) 为基线，n(t) 为高斯白噪声"
            "和可选 AR(1) 低频漂移。电压裁剪后经 ADC 量化为 uint16，同时记录饱和掩码。"
            "触发支持正/负阈值、滞回和最小保持时间；死时间支持非延长型/延长型，只影响观测层。"
        ),
    )

    # 3. 噪声分析方法
    _add_heading(doc, "三、噪声分析方法")

    _add_heading(doc, "3.1 脉冲模式（Phase B）", level=2)
    _add_body(doc, "三种经典估计器作用于观测事件时刻表：")
    _add_body(doc, "  Rossi-α（时间差直方图）：p(τ) = B + A·exp(−α|τ|)")
    _add_body(doc, "  Feynman-α（VTM）：Y(T) = Y∞·[1 − (1−e^{−αT})/(αT)]")
    _add_body(doc, "  PSD（Cohn-α）：F(ω) = C·[1 + Y∞·α²/(α² + ω²)]")
    _add_body(
        doc,
        (
            "bin 间相关不确定度使用时间块 bootstrap；报告值取 bootstrap 与朴素拟合误差的保守上界。"
            "死时间偏置使用 Hazama 一阶非延长型 VTM 修正。"
        ),
    )

    _add_heading(doc, "3.2 连续信号（Phase C）", level=2)
    _add_body(doc, "直接在连续电压波形上计算：")
    _add_body(
        doc,
        (
            "  连续 ACF(θ)：含探测器脉冲衰减常数 α_e ≈ 1/τ_d 的五参数模型；"
            "实用中截除 α_e 衰减区（<200 μs），在长滞后段拟合纯 α 指数项。"
        ),
    )
    _add_body(
        doc,
        ("  连续 VTM(T)：电压门均值方差比；长门宽下 α_e 项饱和为常数偏移，拟合纯 f₁(αT) 形状。"),
    )
    _add_body(
        doc,
        (
            "  Wiener 去卷积：d = F⁻¹{F{c}·W}，W = F*{f}/(|F{f}|² + γ·NSR)；"
            "扫描 γ–NSR 稳定域，确定最优正则化范围。"
        ),
    )
    _add_body(
        doc, ("  双探测器 CCF/CTM：两通道共享同链关联，自项/电子噪声独立，自动消除脉冲形状自项。")
    )
    _add_body(
        doc,
        (
            "  可用边界：扫描 (α × 计数率) 网格，绘出脉冲计数/连续信号/去卷积三条可用边界，"
            "量化 He-3 壁效应 w(η) 对精度的影响。"
        ),
    )

    # 4. 软件架构
    _add_heading(doc, "四、软件架构")
    _add_body(doc, "项目采用 Python 3.11+、src 布局、pyproject.toml 工程化：")
    _add_table(
        doc,
        ["模块", "职责"],
        [
            ["physics/", "泊松/分支链到达、He-3 能谱、幅值、脉冲参数、真值事件"],
            ["synthesis/", "连续波形渲染（直接稀疏/递推）、噪声、基线、ADC"],
            ["acquisition/", "阈值触发、死时间、事件关联表"],
            ["analysis/", "统计验证、脉冲特征、噪声估计（脉冲+连续）、报告与绘图"],
            ["io/", "HDF5 读写、示波器/MCNP/DT5800 stub"],
            ["ui/", "Streamlit 本地 Web 交互界面"],
            ["calibration/", "只读数据资格审查"],
        ],
    )

    # 5. 阶段交付状态
    _add_heading(doc, "五、阶段交付状态")
    _add_table(
        doc,
        ["阶段", "状态", "说明"],
        [
            ["0–3", "human_review_passed", "工程骨架、事件/能谱、连续波形、触发/死时间/数据集"],
            ["A", "human_review_passed", "纯瞬发相关事件、lineage 旁表、6 图验证报告"],
            ["B", "human_review_passed", "三法 α 复原闭合、bootstrap、死时间偏置、7 图报告"],
            ["C", "awaiting_human_review", "连续 ACF/VTM/去卷积、可用边界、壁效应敏感性"],
            ["HIL", "skipped", "DT5800 硬件在环（用户跳过）"],
            ["4–6S", "archived/deferred", "示波器标定、ML 路线已归档"],
        ],
    )

    # 6. 数据格式
    _add_heading(doc, "六、数据格式")
    _add_body(doc, "主格式为 HDF5（分块压缩），结构如下：")
    _add_body(doc, "  /metadata — 配置 JSON、版本、参数状态、随机 seed")
    _add_body(doc, "  /events/true — 真值事件表（event_id, t_s, energy_dep_keV, ...）")
    _add_body(doc, "  /events/observed — 观测事件表（trigger_id, accepted, ...）")
    _add_body(doc, "  /events/lineage — 裂变链旁表（event_id, chain_id, generation）")
    _add_body(doc, "  /blocks — 分块模拟电压、ADC 样本、索引")
    _add_body(doc, "  /windows — 固定长度事件中心 ADC 窗口")
    _add_body(doc, "  /statistics — 计数率、触发、堆积、死时间损失统计")
    _add_body(doc, "所有物理量使用 SI 单位（keV 例外），字段名显式带单位。")

    # 7. 参数真实性
    _add_heading(doc, "七、参数真实性约定")
    _add_body(doc, "所有参数必须标记为三种状态之一：")
    _add_body(doc, "  synthetic_demo — 仅用于演示软件可运行，非实测参数")
    _add_body(doc, "  provisional — 由有限数据或文献估计，等待人工确认")
    _add_body(doc, "  validated — 有明确实验依据且已人工确认")
    _add_body(doc, "当前所有演示值均标记为 synthetic_demo，不声称代表真实探测器、前放或反应堆。")

    # 8. 参考文献
    _add_heading(doc, "八、参考文献")
    refs = [
        "Pázsit I., Pál L. (2008). Neutron Fluctuations. Elsevier.",
        "Williams M.M.R. (1974). Random Processes in Nuclear Reactors. Pergamon.",
        "Boros M.I. 等 (2026). arXiv:2606.10950. (连续信号噪声可行性)",
        "Endo T., Yamamoto A. (2019). Ann. Nucl. Energy 124:606. (Feynman-α bootstrap)",
        "Hazama T. (2003). Ann. Nucl. Energy 30:615. (死时间修正)",
        "Pál L., Pázsit I. (2015). NIMA 794:90. (增殖介质连续信号)",
    ]
    for r in refs:
        _add_body(doc, f"  • {r}")

    out = Path("docs/技术方案.docx")
    doc.save(str(out))
    return out


# ============================================================
# 使用说明书
# ============================================================


def generate_user_manual():
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("He-3 脉冲信号模拟系统 使用说明书", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("软件操作、配置指南、功能概述").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"版本 0.1.0 | {datetime.date.today()}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    # 1. 安装
    _add_heading(doc, "一、环境安装")
    _add_body(doc, "推荐使用 Conda 环境（Python 3.11+）：")
    _add_body(doc, "  conda create -n signal_create python=3.11")
    _add_body(doc, "  conda activate signal_create")
    _add_body(doc, "  cd he3_codex_context")
    _add_body(doc, '  pip install -e ".[dev,ui,doc]"')
    _add_body(doc, "")
    _add_body(
        doc,
        "依赖包：numpy, scipy, matplotlib, h5py, pydantic, PyYAML, typer, streamlit, python-docx",
    )

    # 2. 快速启动
    _add_heading(doc, "二、快速启动")
    _add_body(doc, "交互界面（推荐）：")
    _add_body(doc, "  he3sim ui")
    _add_body(doc, "  → 浏览器打开 http://127.0.0.1:8501")
    _add_body(doc, "")
    _add_body(doc, "命令行方式：")
    _add_body(doc, "  he3sim --help    # 查看所有命令")
    _add_body(doc, "  he3sim validate-config -c configs/demo_minimal.yaml")

    # 3. CLI 命令
    _add_heading(doc, "三、命令行命令一览")
    _add_table(
        doc,
        ["命令", "功能", "所属阶段"],
        [
            ["validate-config", "校验 YAML 配置", "基础"],
            ["simulate-events", "生成真值事件 → HDF5", "Phase 1"],
            ["validate-arrivals", "泊松到达统计验证", "Phase 1"],
            ["validate-correlated", "相关事件静态验证报告", "Phase A"],
            ["simulate-waveform", "生成连续波形 → HDF5", "Phase 2"],
            ["plot-waveform", "波形图片（只读）", "Phase 2"],
            ["generate-dataset", "完整数据集（触发+死时间）", "Phase 3"],
            ["validate-physics", "死时间关系验证", "Phase 3"],
            ["analyze-noise", "单工况 Phase B 噪声分析", "Phase B"],
            ["validate-alpha-recovery", "多工况 α 复原闭环", "Phase B"],
            ["analyze-continuous-noise", "连续 ACF/VTM/去卷积", "Phase C"],
            ["scan-usability-frontier", "可用边界扫描", "Phase C"],
            ["qualify-acquisition", "采集数据只读资格审查", "Phase 4Q"],
            ["ui", "启动交互界面", "UI"],
        ],
    )

    # 4. UI 各页面
    _add_heading(doc, "四、交互界面操作")
    _add_body(doc, "启动 he3sim ui 后，左侧边栏有五个页面：")

    _add_heading(doc, "4.1 事件生成与统计", level=2)
    _add_body(doc, "选择泊松或相关源模型，设置计数率、观察时长、k_eff、α、探测效率等参数，")
    _add_body(doc, "点击「生成事件」即显示事件时序栅格图和间隔分布。相关模式下同时生成泊松对照。")

    _add_heading(doc, "4.2 连续波形生成", level=2)
    _add_body(doc, "配置源模型、计数率、时长、采样率和双指数脉冲参数（τ_r、τ_d），")
    _add_body(doc, "生成含噪声和双指数脉冲叠加的模拟前置放大器电压波形。可拖动滑块选择显示区间。")

    _add_heading(doc, "4.3 Phase B 脉冲噪声", level=2)
    _add_body(doc, "输入相关中子参数，一键运行 Rossi-α、Feynman-α 和 PSD 三种估计器，")
    _add_body(doc, "显示 α 估计值、相对误差和拟合曲线图。支持时间块 bootstrap 不确定度。")

    _add_heading(doc, "4.4 Phase C 连续信号", level=2)
    _add_body(doc, "从连续电压波形直接计算 ACF、VTM 和 Wiener 去卷积，显示 α 复原结果")
    _add_body(doc, "和去卷积三面板对比图。可调节 Wiener 正则化参数 γ。")

    _add_heading(doc, "4.5 历史结果浏览", level=2)
    _add_body(doc, "浏览 outputs/ 下已有的分析报告目录，查看 JSON 指标和 PNG 图表。")

    # 5. 配置文件
    _add_heading(doc, "五、配置文件说明")
    _add_body(doc, "主要配置文件位于 configs/ 目录：")
    _add_body(doc, "  demo_minimal.yaml — 非相关泊松演示配置")
    _add_body(doc, "  demo_correlated.yaml — 相关纯瞬发演示配置")
    _add_body(doc, "  provisional_he3.yaml — 待标定结构模板")
    _add_body(doc, "所有参数值均含 value / status / notes 三元组，status 必为 synthetic_demo、")
    _add_body(doc, "provisional 或 validated 之一。")

    # 6. 已知限制
    _add_heading(doc, "六、已知限制")
    _add_body(doc, "• 所有参数为 synthetic_demo，非实测设备标定值")
    _add_body(doc, "• 仅支持一速点模型纯瞬发分支过程，不含延迟中子")
    _add_body(doc, "• 连续 VTM 在低 SNR 下拟合可能退化")
    _add_body(doc, "• CCF/CTM 未经 DT5800 或真实电子学验证")
    _add_body(doc, "• 不含桌面 GUI、远程 Web 服务、DT5800 实时控制")
    _add_body(doc, "• 当前不接入真实反应堆数据")

    # 7. QA 命令
    _add_heading(doc, "七、质量检查命令")
    _add_body(doc, "修改代码后运行：")
    _add_body(doc, "  python -m ruff format --check .")
    _add_body(doc, "  python -m ruff check .")
    _add_body(doc, "  python -m mypy src")
    _add_body(doc, "  python -m pytest -q")

    out = Path("docs/使用说明书.docx")
    doc.save(str(out))
    return out


if __name__ == "__main__":
    p1 = generate_tech_proposal()
    print(f"技术方案: {p1}")
    p2 = generate_user_manual()
    print(f"使用说明书: {p2}")
